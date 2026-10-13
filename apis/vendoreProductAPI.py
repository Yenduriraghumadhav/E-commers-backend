from typing import Optional
from fastapi import APIRouter , HTTPException , File  , UploadFile, Cookie
from http.client import HTTPException
from starlette import status
from pydantic import BaseModel , Field
from schemas.productSchema import ProductDetailsSchema
from dataBase.dataBaseConnectivity import db
from debaggingFile.loggerFile import logger
from openpyxl import Workbook, load_workbook
from fastapi.responses import StreamingResponse
from io import BytesIO
from uuid import uuid4

router = APIRouter()
vendore_details = db["vendore_collection"]
productDetails = db["product_details"]
jwttoken_toaddProduct = db["VendoreJwtTokenFile"]



@router.post("/vendore_products")
async def create_product(email: str,product_data_input: ProductDetailsSchema,vendor_access_token: str | None = Cookie(default=None)):

    if not vendor_access_token:
        logger.error("JWT token not provided")

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required"
        )

    stored_token = jwttoken_toaddProduct.find_one({
        "token": vendor_access_token
    })

    if not stored_token:
        logger.error("Invalid JWT token")

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid authentication token"
        )

    if not email:
        logger.error("Email not provided")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Email is required"
        )

    product_name = product_data_input.product_name
    category = product_data_input.category
    brand_name = product_data_input.brand_name
    bullet_points = product_data_input.bullet_points
    description = product_data_input.description
    product_mrp = product_data_input.product_mrp
    product_price = product_data_input.product_price
    product_quantity = product_data_input.product_quantity

    if (
            not product_name
            or not category
            or not brand_name
            or not bullet_points
            or not description
            or not product_price
            or not product_quantity
    ):
        logger.error("Enter all product details")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Enter all required product details"
        )

    vendor = vendore_details.find_one({
        "vendore_email": email
    })

    if not vendor:
        logger.error(f"Vendor not found: {email}")

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Vendor not found"
        )

    seller_id = vendor.get("seller_id")

    logger.info(
        f"Vendor seller_id: {seller_id} | Email: {email}"
    )

    if not seller_id:
        logger.warning(
            f"Seller ID not found for vendor: {email}"
        )

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Seller id not found in DB"
        )

    product_id = str(uuid4())

    logger.info(
        f"New product ID generated: {product_id} | Email: {email}"
    )

    product_data = {
        "product_id": product_id,
        "product_name": product_name,
        "brand_name": brand_name,
        "category": category,
        "description": description,
        "bullet_points": bullet_points,
        "product_price": product_price,
        "product_mrp": product_mrp,
        "product_quantity": product_quantity,
        "product_images": [],
        "vendor_email": email,
        "seller_id": seller_id
    }

    result =  productDetails.insert_one(product_data)

    logger.info(
        f"Product created successfully: {result.inserted_id}"
    )

    return {
        "message": "Product created successfully",
        "product_Id": product_id,
        "seller_id": seller_id
    }


class ProductDetailsSchema(BaseModel):
    product_name: str
    brand_name: str
    category: str
    description: str
    bullet_points: list[str]
    product_price: float = Field(gt=0)
    product_mrp: Optional[float] = None
    product_quantity: float = Field(gt=0)


@router.put("/vendore-edit")
async def edit_vendor(sellerids : str , email: str,vendorData: ProductDetailsSchema,jwt_token: str | None = Cookie(default=None)):

    if not jwt_token:
        logger.error("JWT token not provided")

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required"
        )

    if not email:
        logger.error("Email not provided")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Email is required"
        )

    token_verification = await jwttoken_toaddProduct.find_one({
        "vendore_email": email,
        "token": jwt_token
    })

    if not token_verification:
        logger.error(
            f"Vendor edit authentication failed | email={email}"
        )

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Invalid authentication"
        )

    vendor = await vendore_details.find_one({
        "vendore_email": email
    })

    if not vendor:
        logger.error(
            f"Vendor not found | email={email}"
        )

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Vendor not found"
        )

    seller_id = jwttoken_toaddProduct.find({
        "seller_id" : sellerids
    })

    if not seller_id:
        logger.error(
            f"Seller ID not found in DB To the| email={email}"
        )

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Seller ID not found in DB"
        )

    update_data = {
        "product_name": vendorData.product_name,
        "brand_name": vendorData.brand_name,
        "category": vendorData.category,
        "description": vendorData.description,
        "bullet_points": vendorData.bullet_points,
        "product_price": vendorData.product_price,
        "product_mrp": vendorData.product_mrp,
        "product_quantity": vendorData.product_quantity
    }

    result = await productDetails.update_one(
        {
            "seller_id": seller_id
        },
        {
            "$set": update_data
        }
    )

    if result.matched_count == 0:

        logger.warning(
            f"Product not found or ownership mismatch | email={email}"
        )

        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Product not found or you are not the owner"
        )

    if result.modified_count == 0:

        logger.info(
            f"No product changes | email={email}"
        )

        return {
            "message": "No changes made"
        }

    logger.info(
        f"Product updated successfully | email={email}"
    )

    return {
        "message": "Product updated successfully",
    }


@router.get("/vendore-get")
async def get_vendore(userEmail:str , jwtToken : str | None = Cookie(default=None)):

    if not jwtToken:
        logger.error(
            "JWT token not provided"
        )
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required"
        )

    if not userEmail:
        logger.error(
            "Email not provided"
        )
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required"
        )

    filteringVendoreData = jwttoken_toaddProduct.find_one({
        "vendore_email" : userEmail,
        "token" : jwtToken
    })

    if not filteringVendoreData:
        logger.warning(
            "vendoreEmail or Jwt Token not found in DB"
        )
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication Failed | Details Not Found in DB"
        )

    productsFilteringDetails = list(productDetails.find_one({ "vendor_email" : userEmail }))

    if not productsFilteringDetails:
        logger.info(
            f"Product details not found in DB for email | {userEmail}"
        )
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail={
                "message": "No product data found for this email.",
                "suggestion": "Please try adding a product. If the problem continues, please contact the service team."
            }
        )

    for product in productsFilteringDetails:
        product.pop("_id", None)

    return {
        "productDetails":productsFilteringDetails,
        "vendoreEmail" : userEmail
    }


@router.get("/download_vendore_productsExcel_template")
async def download_vendore_productsExcel_template():

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Products"

    headers = [
        "product_name",
        "brand_name",
        "category",
        "description",
        "bullet_points",
        "product_price",
        "product_mrp",
        "product_quantity"
    ]

    worksheet.append(headers)

    worksheet.append([
        "iPhone 15",
        "Apple",
        "Mobile",
        "Apple iPhone 15 smartphone",
        "48MP Camera,5G,OLED Display",
        65000,
        70000,
        10
    ])

    excel_file = BytesIO()

    workbook.save(excel_file)

    excel_file.seek(0)

    return StreamingResponse(
        excel_file,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition":
                "attachment; filename=vendor_products_template.xlsx"
        }
    )


@router.post("/upload_vendore_productsExcel")
async def upload_vendore_productsExcel(vendorEmail: str,file: UploadFile = File(...),vendor_access_token: str | None = Cookie(default=None)):


    if not vendor_access_token:
        logger.error("JWT token not provided")

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required"
        )

    if not vendorEmail:
        logger.error("Vendor email not provided")

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication required"
        )

    filer_Data =  jwttoken_toaddProduct.find_one({
        "token": vendor_access_token
    })

    vendore_Data = vendore_details.find_one({"vendore_email": vendorEmail})

    if not vendore_Data:
        logger.error(
            "vendore email is not found in vendoreCollect File DB"
        )

    seller_id = vendore_Data.get("seller_id")

    if not seller_id:
        logger.warning(
            f"seller_id not get in the DB To the| email={vendorEmail}"
        )

    if not filer_Data:
        logger.error(
            f"Vendor email or JWT token not found | email={vendorEmail}"
        )

        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="Authentication failed | email or JWT token not found"
        )

    if not file.filename:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Excel file is required"
        )

    if not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Only .xlsx Excel files are supported"
        )

    file_content = await file.read()

    try:
        workbook = load_workbook(
            filename=BytesIO(file_content),
            data_only=True
        )

    except Exception as e:
        logger.error(f"Failed to read Excel file: {e}")

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid Excel file"
        )

    worksheet = workbook.active

    headers = [
        str(cell.value).strip()
        if cell.value is not None
        else ""
        for cell in worksheet[1]
    ]

    required_headers = [
        "product_name",
        "brand_name",
        "category",
        "description",
        "bullet_points",
        "product_price",
        "product_mrp",
        "product_quantity"
    ]

    if headers != required_headers:

        logger.error(f"Invalid Excel headers: {headers}" )

        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail={
                "message": "Invalid Excel format",
                "required_headers": required_headers,
                "received_headers": headers
            }
        )

    inserted_products = []
    errors = []

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            values_only=True
        ),
        start=2
    ):

        if not any(value is not None for value in row):
            continue

        try:

            product_name = row[0]
            brand_name = row[1]
            category = row[2]
            description = row[3]
            bullet_points = row[4]
            product_price = row[5]
            product_mrp = row[6]
            product_quantity = row[7]

            if not product_name:
                errors.append({
                    "row": row_number,
                    "error": "product_name is required"
                })
                continue

            if not brand_name:
                errors.append({
                    "row": row_number,
                    "error": "brand_name is required"
                })
                continue

            if not category:
                errors.append({
                    "row": row_number,
                    "error": "category is required"
                })
                continue

            if product_price is None:
                errors.append({
                    "row": row_number,
                    "error": "product_price is required"
                })
                continue

            if product_mrp is None:
                errors.append({
                    "row": row_number,
                    "error": "product_mrp is required"
                })
                continue

            if product_quantity is None:
                errors.append({
                    "row": row_number,
                    "error": "product_quantity is required"
                })
                continue

            product_id = str(uuid4())

            product_data = {
                "product_id": product_id,
                "product_name": product_name,
                "brand_name": brand_name,
                "category": category,
                "description": description,
                "bullet_points": bullet_points,
                "product_price": product_price,
                "product_mrp": product_mrp,
                "product_quantity": product_quantity,
                "vendor_email": vendorEmail,
                "seller_id": seller_id
            }

            productDetails.insert_one( product_data )

            inserted_products.append({
                "row": row_number,
                "product_id": product_id,
                "product_name": product_name
            })

        except Exception as e:

            logger.error(
                f"Error processing Excel row "
                f"{row_number}: {e}"
            )

            errors.append({
                "row": row_number,
                "error": str(e)
            })

    return {
        "success": True,
        "message": "Excel processed successfully",
        "vendor_email": vendorEmail,
        "total_inserted": len(inserted_products),
        "total_errors": len(errors),
        "inserted_products": inserted_products,
        "errors": errors,
        "sellerID" : seller_id
    }


